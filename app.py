# -*- coding: utf-8 -*-
"""Accessibility Tracker — valutazione conformità WCAG 2.2 (A/AA)."""
import csv
import io
import os

from flask import (Flask, flash, redirect, render_template, request,
                   send_file, url_for)

from database import get_db, init_db, unique_slug
from wcag_data import WCAG_BASE_URL

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "accessibility-tracker-local")

STATUS_LABELS = {
    "da_verificare": "Da verificare",
    "conforme": "Conforme",
    "non_conforme": "Non conforme",
    "non_applicabile": "Non applicabile",
}
WORK_LABELS = {
    "": "—",
    "aperta": "Aperta",
    "in_lavorazione": "In lavorazione",
    "risolta": "Risolta",
    "verificata": "Risolta e verificata",
}
PROJECT_STATUSES = ["In corso", "In attesa", "Completato"]


@app.context_processor
def inject_globals():
    return {"STATUS_LABELS": STATUS_LABELS, "WORK_LABELS": WORK_LABELS,
            "PROJECT_STATUSES": PROJECT_STATUSES, "WCAG_BASE_URL": WCAG_BASE_URL}


STATS_SELECT = """
    SELECT
      COUNT(e.id) AS totale,
      SUM(e.status = 'conforme') AS conformi,
      SUM(e.status = 'non_conforme') AS non_conformi,
      SUM(e.status = 'da_verificare') AS da_verificare,
      SUM(e.status = 'non_applicabile') AS non_applicabili,
      SUM(e.status = 'non_conforme' AND e.work_status IN ('risolta', 'verificata')) AS risolte,
      SUM(e.status = 'non_conforme' AND cr.level = 'AAA') AS non_conformi_aaa
    FROM evaluations e
    JOIN criteria cr ON cr.id = e.criterion_id
    JOIN components c ON c.id = e.component_id
    JOIN pages p ON p.id = c.page_id
    JOIN projects pr ON pr.id = p.project_id
"""

# Conformità per criterio distinto sul totale dei criteri A/AA esistenti (non sulle
# rilevazioni): un criterio conta una sola volta per l'intero progetto/piattaforma
# anche se violato in più componenti, e i criteri mai rilevati si assumono conformi.
CRITERIA_NON_CONFORMI_SELECT = """
    SELECT COUNT(DISTINCT cr.id) AS n
    FROM evaluations e
    JOIN criteria cr ON cr.id = e.criterion_id
    JOIN components c ON c.id = e.component_id
    JOIN pages p ON p.id = c.page_id
    JOIN projects pr ON pr.id = p.project_id
    WHERE cr.level != 'AAA' AND e.status = 'non_conforme' AND {where}
"""


def _finish_stats(row):
    s = dict(row)
    for k in s:
        s[k] = s[k] or 0
    s["pct_avanzamento"] = round(100 * (s["totale"] - s["da_verificare"]) / s["totale"]) if s["totale"] else 0
    return s


def _pct_conformi(db, where, params):
    totale_aa = db.execute("SELECT COUNT(*) FROM criteria WHERE level != 'AAA'").fetchone()[0]
    non_conformi = db.execute(CRITERIA_NON_CONFORMI_SELECT.format(where=where), params).fetchone()["n"]
    return round(100 * (totale_aa - non_conformi) / totale_aa) if totale_aa else 0


def project_stats(db, project_id):
    s = _finish_stats(db.execute(STATS_SELECT + " WHERE p.project_id = ?", (project_id,)).fetchone())
    s["pct_conformi"] = _pct_conformi(db, "p.project_id = ?", (project_id,))
    return s


def platform_stats(db, platform_id):
    s = _finish_stats(db.execute(STATS_SELECT + " WHERE pr.platform_id = ?", (platform_id,)).fetchone())
    s["pct_conformi"] = _pct_conformi(db, "pr.platform_id = ?", (platform_id,))
    return s


# ---------------------------------------------------------------- Dashboard
def _get_platform(db, slug):
    return db.execute("SELECT * FROM platforms WHERE slug = ?", (slug,)).fetchone()


def _get_project(db, platform_slug, project_slug):
    return db.execute("""SELECT p.*, pl.name AS platform_name, pl.slug AS platform_slug
                         FROM projects p JOIN platforms pl ON pl.id = p.platform_id
                         WHERE pl.slug = ? AND p.slug = ?""",
                      (platform_slug, project_slug)).fetchone()


@app.route("/")
def index():
    db = get_db()
    platforms = db.execute("SELECT * FROM platforms ORDER BY name").fetchall()
    agg = {pl["id"]: platform_stats(db, pl["id"]) for pl in platforms}
    counts = {pl["id"]: db.execute(
        "SELECT COUNT(*) FROM projects WHERE platform_id = ?", (pl["id"],)).fetchone()[0]
        for pl in platforms}
    db.close()
    return render_template("index.html", platforms=platforms, agg=agg,
                           project_counts=counts)


# ---------------------------------------------------------------- Piattaforme
@app.route("/platform", methods=["POST"])
def new_platform():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Il nome della piattaforma è obbligatorio.", "error")
        return redirect(url_for("index"))
    db = get_db()
    slug = unique_slug(db, "platforms", name, fallback="piattaforma")
    db.execute("INSERT INTO platforms (name, slug, description) VALUES (?, ?, ?)",
               (name, slug, request.form.get("description", "").strip()))
    db.commit()
    db.close()
    flash(f"Piattaforma «{name}» creata. Ora crea il suo primo progetto.", "ok")
    return redirect(url_for("platform_detail", platform_slug=slug))


@app.route("/platform/<platform_slug>")
def platform_detail(platform_slug):
    db = get_db()
    platform = _get_platform(db, platform_slug)
    if not platform:
        db.close()
        flash("Piattaforma non trovata.", "error")
        return redirect(url_for("index"))
    projects = db.execute("SELECT * FROM projects WHERE platform_id = ? ORDER BY created_at",
                          (platform["id"],)).fetchall()
    stats = {p["id"]: project_stats(db, p["id"]) for p in projects}
    agg = platform_stats(db, platform["id"])
    db.close()
    return render_template("platform.html", platform=platform, projects=projects,
                           stats=stats, agg=agg)


@app.route("/platform/<platform_slug>/edit", methods=["POST"])
def edit_platform(platform_slug):
    db = get_db()
    platform = _get_platform(db, platform_slug)
    if not platform:
        db.close()
        flash("Piattaforma non trovata.", "error")
        return redirect(url_for("index"))
    name = request.form.get("name", "").strip() or "Senza nome"
    slug = platform["slug"]
    if name != platform["name"]:
        slug = unique_slug(db, "platforms", name, exclude_id=platform["id"], fallback="piattaforma")
    db.execute("UPDATE platforms SET name = ?, slug = ?, description = ? WHERE id = ?",
               (name, slug, request.form.get("description", "").strip(), platform["id"]))
    db.commit()
    db.close()
    flash("Piattaforma aggiornata.", "ok")
    return redirect(url_for("platform_detail", platform_slug=slug))


@app.route("/platform/<platform_slug>/delete", methods=["POST"])
def delete_platform(platform_slug):
    db = get_db()
    platform = _get_platform(db, platform_slug)
    if not platform:
        db.close()
        flash("Piattaforma non trovata.", "error")
        return redirect(url_for("index"))
    n = db.execute("SELECT COUNT(*) FROM projects WHERE platform_id = ?", (platform["id"],)).fetchone()[0]
    if n:
        db.close()
        flash(f"Impossibile eliminare: la piattaforma contiene {n} progetti. "
              "Elimina o sposta prima i progetti.", "error")
        return redirect(url_for("platform_detail", platform_slug=platform_slug))
    db.execute("DELETE FROM platforms WHERE id = ?", (platform["id"],))
    db.commit()
    db.close()
    flash("Piattaforma eliminata.", "ok")
    return redirect(url_for("index"))


@app.route("/platform/<platform_slug>/project", methods=["POST"])
def new_project(platform_slug):
    db = get_db()
    platform = _get_platform(db, platform_slug)
    if not platform:
        db.close()
        flash("Piattaforma non trovata.", "error")
        return redirect(url_for("index"))
    name = request.form.get("name", "").strip()
    if not name:
        db.close()
        flash("Il nome del progetto è obbligatorio.", "error")
        return redirect(url_for("platform_detail", platform_slug=platform_slug))
    slug = unique_slug(db, "projects", name, "platform_id", platform["id"], fallback="progetto")
    db.execute(
        "INSERT INTO projects (name, slug, description, url, platform_id) VALUES (?, ?, ?, ?, ?)",
        (name, slug, request.form.get("description", "").strip(),
         request.form.get("url", "").strip(), platform["id"]))
    db.commit()
    db.close()
    flash(f"Progetto «{name}» creato.", "ok")
    return redirect(url_for("project_detail", platform_slug=platform_slug, project_slug=slug))


@app.route("/platform/<platform_slug>/project/<project_slug>/edit", methods=["POST"])
def edit_project(platform_slug, project_slug):
    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if not project:
        db.close()
        flash("Progetto non trovato.", "error")
        return redirect(url_for("index"))
    name = request.form.get("name", "").strip() or "Senza nome"
    new_platform_id = int(request.form.get("platform_id") or project["platform_id"])
    slug = project["slug"]
    if name != project["name"] or new_platform_id != project["platform_id"]:
        slug = unique_slug(db, "projects", name, "platform_id", new_platform_id,
                           exclude_id=project["id"], fallback="progetto")
    db.execute(
        "UPDATE projects SET name = ?, slug = ?, description = ?, url = ?, status = ?, platform_id = ? WHERE id = ?",
        (name, slug, request.form.get("description", "").strip(),
         request.form.get("url", "").strip(),
         request.form.get("status", "In corso"), new_platform_id, project["id"]))
    db.commit()
    pl_slug = db.execute("SELECT slug FROM platforms WHERE id = ?", (new_platform_id,)).fetchone()["slug"]
    db.close()
    flash("Progetto aggiornato.", "ok")
    return redirect(url_for("project_detail", platform_slug=pl_slug, project_slug=slug))


@app.route("/platform/<platform_slug>/project/<project_slug>/delete", methods=["POST"])
def delete_project(platform_slug, project_slug):
    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if project:
        db.execute("DELETE FROM projects WHERE id = ?", (project["id"],))
        db.commit()
        flash("Progetto eliminato.", "ok")
    db.close()
    return redirect(url_for("platform_detail", platform_slug=platform_slug))


# ---------------------------------------------------------------- Progetto
@app.route("/platform/<platform_slug>/project/<project_slug>")
def project_detail(platform_slug, project_slug):
    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if not project:
        db.close()
        flash("Progetto non trovato.", "error")
        return redirect(url_for("index"))
    platforms = db.execute("SELECT * FROM platforms ORDER BY name").fetchall()
    pages = db.execute("""
        SELECT p.*,
          (SELECT COUNT(*) FROM components c WHERE c.page_id = p.id) AS n_componenti,
          (SELECT COUNT(*) FROM evaluations e JOIN components c ON c.id = e.component_id
             WHERE c.page_id = p.id AND e.status = 'non_conforme') AS n_anomalie,
          (SELECT COUNT(*) FROM evaluations e JOIN components c ON c.id = e.component_id
             WHERE c.page_id = p.id AND e.status = 'da_verificare') AS n_da_verificare
        FROM pages p WHERE p.project_id = ? ORDER BY p.created_at""", (project["id"],)).fetchall()
    stats = project_stats(db, project["id"])
    db.close()
    return render_template("project.html", project=project, pages=pages,
                           stats=stats, platforms=platforms)


@app.route("/platform/<platform_slug>/project/<project_slug>/pages/new", methods=["POST"])
def new_page(platform_slug, project_slug):
    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if not project:
        db.close()
        flash("Progetto non trovato.", "error")
        return redirect(url_for("index"))
    name = request.form.get("name", "").strip()
    if name:
        db.execute("INSERT INTO pages (project_id, name, url, notes) VALUES (?, ?, ?, ?)",
                   (project["id"], name, request.form.get("url", "").strip(),
                    request.form.get("notes", "").strip()))
        db.commit()
        flash(f"Pagina «{name}» aggiunta.", "ok")
    else:
        flash("Il nome della pagina è obbligatorio.", "error")
    db.close()
    return redirect(url_for("project_detail", platform_slug=platform_slug, project_slug=project_slug))


@app.route("/pagine/<int:page_id>/elimina", methods=["POST"])
def delete_page(page_id):
    db = get_db()
    row = db.execute("""SELECT pr.slug AS project_slug, pl.slug AS platform_slug
                        FROM pages p JOIN projects pr ON pr.id = p.project_id
                        JOIN platforms pl ON pl.id = pr.platform_id
                        WHERE p.id = ?""", (page_id,)).fetchone()
    db.execute("DELETE FROM pages WHERE id = ?", (page_id,))
    db.commit()
    db.close()
    flash("Pagina eliminata.", "ok")
    return redirect(url_for("project_detail", platform_slug=row["platform_slug"],
                            project_slug=row["project_slug"]) if row else url_for("index"))


# ---------------------------------------------------------------- Pagina
@app.route("/pagine/<int:page_id>")
def page_detail(page_id):
    db = get_db()
    page = db.execute("""SELECT p.*, pr.name AS project_name, pr.id AS project_id,
                                pr.slug AS project_slug, pl.slug AS platform_slug
                         FROM pages p JOIN projects pr ON pr.id = p.project_id
                         JOIN platforms pl ON pl.id = pr.platform_id
                         WHERE p.id = ?""", (page_id,)).fetchone()
    if not page:
        db.close()
        flash("Pagina non trovata.", "error")
        return redirect(url_for("index"))
    components = db.execute("""
        SELECT c.*, t.name AS type_name,
          (SELECT COUNT(*) FROM evaluations e WHERE e.component_id = c.id) AS n_criteri,
          (SELECT COUNT(*) FROM evaluations e WHERE e.component_id = c.id AND e.status = 'conforme') AS n_conformi,
          (SELECT COUNT(*) FROM evaluations e WHERE e.component_id = c.id AND e.status = 'non_conforme') AS n_anomalie,
          (SELECT COUNT(*) FROM evaluations e WHERE e.component_id = c.id AND e.status = 'da_verificare') AS n_da_verificare
        FROM components c JOIN component_types t ON t.id = c.component_type_id
        WHERE c.page_id = ? ORDER BY c.created_at""", (page_id,)).fetchall()
    types = db.execute("SELECT * FROM component_types ORDER BY name").fetchall()
    db.close()
    return render_template("page.html", page=page, components=components, types=types)


@app.route("/pagine/<int:page_id>/componenti/nuovo", methods=["POST"])
def new_component(page_id):
    name = request.form.get("name", "").strip()
    type_id = request.form.get("component_type_id")
    if not name or not type_id:
        flash("Nome e tipologia del componente sono obbligatori.", "error")
        return redirect(url_for("page_detail", page_id=page_id))
    db = get_db()
    cur = db.execute(
        "INSERT INTO components (page_id, component_type_id, name, location, notes) VALUES (?, ?, ?, ?, ?)",
        (page_id, type_id, name, request.form.get("location", "").strip(),
         request.form.get("notes", "").strip()))
    comp_id = cur.lastrowid
    # Genera automaticamente le valutazioni per i criteri applicabili alla tipologia
    include_aaa = 1 if request.form.get("include_aaa") else 0
    db.execute("""INSERT INTO evaluations (component_id, criterion_id)
                  SELECT ?, ctc.criterion_id FROM component_type_criteria ctc
                  JOIN criteria cr ON cr.id = ctc.criterion_id
                  WHERE ctc.component_type_id = ? AND (? = 1 OR cr.level != 'AAA')""",
               (comp_id, type_id, include_aaa))
    db.commit()
    db.close()
    flash(f"Componente «{name}» aggiunto con la checklist dei criteri applicabili.", "ok")
    return redirect(url_for("component_detail", component_id=comp_id))


@app.route("/componenti/<int:component_id>/elimina", methods=["POST"])
def delete_component(component_id):
    db = get_db()
    row = db.execute("SELECT page_id FROM components WHERE id = ?", (component_id,)).fetchone()
    db.execute("DELETE FROM components WHERE id = ?", (component_id,))
    db.commit()
    db.close()
    flash("Componente eliminato.", "ok")
    return redirect(url_for("page_detail", page_id=row["page_id"]) if row else url_for("index"))


# ---------------------------------------------------------------- Componente / checklist
@app.route("/componenti/<int:component_id>")
def component_detail(component_id):
    db = get_db()
    comp = db.execute("""
        SELECT c.*, t.name AS type_name, t.aria_notes,
               p.name AS page_name, p.id AS page_id,
               pr.name AS project_name, pr.id AS project_id,
               pr.slug AS project_slug, pl.slug AS platform_slug
        FROM components c
        JOIN component_types t ON t.id = c.component_type_id
        JOIN pages p ON p.id = c.page_id
        JOIN projects pr ON pr.id = p.project_id
        JOIN platforms pl ON pl.id = pr.platform_id
        WHERE c.id = ?""", (component_id,)).fetchone()
    if not comp:
        db.close()
        flash("Componente non trovato.", "error")
        return redirect(url_for("index"))
    evals = db.execute("""
        SELECT e.*, cr.code, cr.title, cr.level, cr.principle, cr.anchor
        FROM evaluations e JOIN criteria cr ON cr.id = e.criterion_id
        WHERE e.component_id = ? ORDER BY cr.sort_order""", (component_id,)).fetchall()
    remaining = db.execute("""
        SELECT * FROM criteria WHERE id NOT IN
          (SELECT criterion_id FROM evaluations WHERE component_id = ?)
        ORDER BY sort_order""", (component_id,)).fetchall()
    hist = db.execute("""
        SELECT h.*, cr.code FROM history h
        JOIN evaluations e ON e.id = h.evaluation_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE e.component_id = ? ORDER BY h.changed_at DESC LIMIT 20""", (component_id,)).fetchall()
    db.close()
    return render_template("component.html", comp=comp, evals=evals,
                           remaining=remaining, history=hist)


@app.route("/componenti/<int:component_id>/valutazioni", methods=["POST"])
def save_evaluations(component_id):
    db = get_db()
    evals = db.execute("SELECT * FROM evaluations WHERE component_id = ?", (component_id,)).fetchall()
    changed = 0
    for e in evals:
        new_status = request.form.get(f"status_{e['id']}", e["status"])
        new_work = request.form.get(f"work_{e['id']}", e["work_status"])
        new_note = request.form.get(f"note_{e['id']}", e["note"]).strip()
        if new_status != "non_conforme":
            new_work = ""
        elif new_status == "non_conforme" and not new_work:
            new_work = "aperta"
        if (new_status, new_work, new_note) != (e["status"], e["work_status"], e["note"]):
            if (new_status, new_work) != (e["status"], e["work_status"]):
                db.execute("""INSERT INTO history (evaluation_id, old_status, new_status,
                              old_work_status, new_work_status, note)
                              VALUES (?, ?, ?, ?, ?, ?)""",
                           (e["id"], e["status"], new_status, e["work_status"], new_work, new_note))
            db.execute("""UPDATE evaluations SET status = ?, work_status = ?, note = ?,
                          updated_at = datetime('now', 'localtime') WHERE id = ?""",
                       (new_status, new_work, new_note, e["id"]))
            changed += 1
    db.commit()
    db.close()
    flash(f"Valutazioni salvate ({changed} aggiornate)." if changed else "Nessuna modifica da salvare.", "ok")
    return redirect(url_for("component_detail", component_id=component_id))


@app.route("/componenti/<int:component_id>/aggiungi-criterio", methods=["POST"])
def add_criterion(component_id):
    crit_id = request.form.get("criterion_id")
    if crit_id:
        db = get_db()
        db.execute("INSERT OR IGNORE INTO evaluations (component_id, criterion_id) VALUES (?, ?)",
                   (component_id, crit_id))
        db.commit()
        db.close()
        flash("Criterio aggiunto alla checklist.", "ok")
    return redirect(url_for("component_detail", component_id=component_id))


@app.route("/valutazioni/<int:eval_id>/rimuovi", methods=["POST"])
def remove_evaluation(eval_id):
    db = get_db()
    row = db.execute("SELECT component_id FROM evaluations WHERE id = ?", (eval_id,)).fetchone()
    db.execute("DELETE FROM evaluations WHERE id = ?", (eval_id,))
    db.commit()
    db.close()
    flash("Criterio rimosso dalla checklist.", "ok")
    return redirect(url_for("component_detail", component_id=row["component_id"]) if row else url_for("index"))


# ---------------------------------------------------------------- Tipologie
@app.route("/tipologie")
def types_list():
    db = get_db()
    types = db.execute("""
        SELECT t.*, (SELECT COUNT(*) FROM component_type_criteria x
                     WHERE x.component_type_id = t.id) AS n_criteri
        FROM component_types t ORDER BY t.name""").fetchall()
    db.close()
    return render_template("tipologie.html", types=types)


@app.route("/tipologie/nuova", methods=["GET", "POST"])
@app.route("/tipologie/<int:type_id>", methods=["GET", "POST"])
def type_form(type_id=None):
    db = get_db()
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("Il nome della tipologia è obbligatorio.", "error")
            db.close()
            return redirect(url_for("types_list"))
        desc = request.form.get("description", "").strip()
        aria = request.form.get("aria_notes", "").strip()
        if type_id is None:
            cur = db.execute("INSERT INTO component_types (name, description, aria_notes) VALUES (?, ?, ?)",
                             (name, desc, aria))
            type_id = cur.lastrowid
        else:
            db.execute("UPDATE component_types SET name = ?, description = ?, aria_notes = ? WHERE id = ?",
                       (name, desc, aria, type_id))
            db.execute("DELETE FROM component_type_criteria WHERE component_type_id = ?", (type_id,))
        for cid in request.form.getlist("criteria"):
            db.execute("INSERT OR IGNORE INTO component_type_criteria VALUES (?, ?)", (type_id, cid))
        db.commit()
        db.close()
        flash("Tipologia salvata.", "ok")
        return redirect(url_for("types_list"))

    ctype = db.execute("SELECT * FROM component_types WHERE id = ?", (type_id,)).fetchone() if type_id else None
    selected = {r["criterion_id"] for r in db.execute(
        "SELECT criterion_id FROM component_type_criteria WHERE component_type_id = ?",
        (type_id,)).fetchall()} if type_id else set()
    criteria = db.execute("SELECT * FROM criteria ORDER BY sort_order").fetchall()
    db.close()
    return render_template("tipologia_form.html", ctype=ctype, criteria=criteria, selected=selected)


@app.route("/criteri")
def criteria_list():
    db = get_db()
    livello = request.args.get("livello", "").upper()
    if livello not in ("A", "AA", "AAA"):
        livello = ""
    query = "SELECT * FROM criteria"
    params = ()
    if livello:
        query += " WHERE level = ?"
        params = (livello,)
    criteria = db.execute(query + " ORDER BY sort_order", params).fetchall()
    counts = {r["level"]: r["n"] for r in
             db.execute("SELECT level, COUNT(*) AS n FROM criteria GROUP BY level").fetchall()}
    db.close()
    return render_template("criteri.html", criteria=criteria, livello=livello, counts=counts)


@app.route("/tipologie/<int:type_id>/elimina", methods=["POST"])
def delete_type(type_id):
    db = get_db()
    used = db.execute("SELECT COUNT(*) FROM components WHERE component_type_id = ?", (type_id,)).fetchone()[0]
    if used:
        flash("Impossibile eliminare: la tipologia è usata da componenti esistenti.", "error")
    else:
        db.execute("DELETE FROM component_types WHERE id = ?", (type_id,))
        db.commit()
        flash("Tipologia eliminata.", "ok")
    db.close()
    return redirect(url_for("types_list"))


# ---------------------------------------------------------------- Report
def report_data(db, project_id):
    """Anomalie (non conformità) raggruppate per pagina e componente."""
    rows = db.execute("""
        SELECT pg.id AS page_id, pg.name AS page_name, pg.url AS page_url,
               c.id AS component_id, c.name AS component_name, c.location,
               t.name AS type_name,
               cr.code, cr.title, cr.level, cr.anchor,
               e.id AS eval_id, e.status, e.work_status, e.note, e.updated_at
        FROM evaluations e
        JOIN components c ON c.id = e.component_id
        JOIN component_types t ON t.id = c.component_type_id
        JOIN pages pg ON pg.id = c.page_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE pg.project_id = ? AND e.status = 'non_conforme'
        ORDER BY pg.created_at, c.created_at, cr.sort_order""", (project_id,)).fetchall()
    return rows


@app.route("/platform/<platform_slug>/project/<project_slug>/report")
def report(platform_slug, project_slug):
    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if not project:
        db.close()
        return redirect(url_for("index"))
    project_id = project["id"]
    anomalies = report_data(db, project_id)
    stats = project_stats(db, project_id)
    histories = {}
    for a in anomalies:
        histories[a["eval_id"]] = db.execute(
            "SELECT * FROM history WHERE evaluation_id = ? ORDER BY changed_at", (a["eval_id"],)).fetchall()
    # criteri violati distinti
    violated = db.execute("""
        SELECT DISTINCT cr.code, cr.title, cr.level, cr.sort_order
        FROM evaluations e
        JOIN components c ON c.id = e.component_id
        JOIN pages pg ON pg.id = c.page_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE pg.project_id = ? AND e.status = 'non_conforme' ORDER BY cr.sort_order""", (project_id,)).fetchall()
    db.close()
    from datetime import datetime
    return render_template("report.html", project=project, anomalies=anomalies,
                           stats=stats, histories=histories, violated=violated,
                           generated=datetime.now().strftime("%d/%m/%Y %H:%M"))


@app.route("/platform/<platform_slug>/project/<project_slug>/export.csv")
def export_csv(platform_slug, project_slug):
    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if not project:
        db.close()
        return redirect(url_for("index"))
    anomalies = report_data(db, project["id"])
    db.close()
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Pagina", "URL pagina", "Componente", "Tipologia", "Posizione",
                "Criterio WCAG", "Titolo criterio", "Livello", "Stato lavorazione",
                "Motivazione / Note", "Ultimo aggiornamento"])
    for a in anomalies:
        w.writerow([a["page_name"], a["page_url"], a["component_name"], a["type_name"],
                    a["location"], a["code"], a["title"], a["level"],
                    WORK_LABELS.get(a["work_status"], a["work_status"]),
                    a["note"], a["updated_at"]])
    data = io.BytesIO(("﻿" + out.getvalue()).encode("utf-8"))
    fname = f"anomalie_{project['name'].replace(' ', '_')}.csv" if project else "anomalie.csv"
    return send_file(data, mimetype="text/csv", as_attachment=True, download_name=fname)


@app.route("/platform/<platform_slug>/project/<project_slug>/export.xlsx")
def export_xlsx(platform_slug, project_slug):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    db = get_db()
    project = _get_project(db, platform_slug, project_slug)
    if not project:
        db.close()
        return redirect(url_for("index"))
    project_id = project["id"]
    anomalies = report_data(db, project_id)
    all_rows = db.execute("""
        SELECT pg.name AS page_name, c.name AS component_name, t.name AS type_name,
               cr.code, cr.title, cr.level, e.status, e.work_status, e.note, e.updated_at
        FROM evaluations e
        JOIN components c ON c.id = e.component_id
        JOIN component_types t ON t.id = c.component_type_id
        JOIN pages pg ON pg.id = c.page_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE pg.project_id = ?
        ORDER BY pg.created_at, c.created_at, cr.sort_order""", (project_id,)).fetchall()
    db.close()

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")

    def sheet(ws, headers, rows):
        ws.append(headers)
        for i in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=i)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        widths = [len(h) for h in headers]
        for r in rows:
            for i, v in enumerate(r):
                widths[i] = min(60, max(widths[i], len(str(v or ""))))
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt + 2
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws1 = wb.active
    ws1.title = "Anomalie"
    sheet(ws1,
          ["Pagina", "Componente", "Tipologia", "Posizione", "Criterio", "Titolo",
           "Livello", "Stato lavorazione", "Motivazione / Note", "Aggiornato il"],
          [[a["page_name"], a["component_name"], a["type_name"], a["location"],
            a["code"], a["title"], a["level"],
            WORK_LABELS.get(a["work_status"], a["work_status"]), a["note"], a["updated_at"]]
           for a in anomalies])

    ws2 = wb.create_sheet("Tutte le valutazioni")
    sheet(ws2,
          ["Pagina", "Componente", "Tipologia", "Criterio", "Titolo", "Livello",
           "Esito", "Stato lavorazione", "Note", "Aggiornato il"],
          [[r["page_name"], r["component_name"], r["type_name"], r["code"], r["title"],
            r["level"], STATUS_LABELS.get(r["status"], r["status"]),
            WORK_LABELS.get(r["work_status"], r["work_status"]), r["note"], r["updated_at"]]
           for r in all_rows])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"report_{project['name'].replace(' ', '_')}.xlsx" if project else "report.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    init_db()
    app.run(host=os.environ.get("FLASK_HOST", "127.0.0.1"),
            port=int(os.environ.get("FLASK_PORT", "8000")),
            debug=False)
