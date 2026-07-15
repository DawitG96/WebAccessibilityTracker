# -*- coding: utf-8 -*-
"""Estrazione dati e generazione file di report (CSV/XLSX) per le anomalie di progetto."""
import csv
import io

from extensions import STATUS_LABELS, WORK_LABELS


def report_data(db, project_id):
    """Anomalie (non conformità) raggruppate per pagina e componente."""
    rows = db.execute("""
        SELECT pg.id AS page_id, pg.name AS page_name, pg.url AS page_url,
               c.id AS component_id, c.name AS component_name, c.location,
               t.name AS type_name,
               cr.code, cr.title, cr.level, cr.anchor,
               e.id AS eval_id, e.status, e.work_status, e.note, e.updated_at
        FROM evaluations e
        JOIN components c ON c.id = e.component_id
        JOIN component_types t ON t.id = c.component_type_id
        JOIN pages pg ON pg.id = c.page_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE pg.project_id = ? AND e.status = 'non_conforme'
        ORDER BY pg.created_at, c.created_at, cr.sort_order""", (project_id,)).fetchall()
    return rows


def evaluation_histories(db, anomalies):
    return {
        a["eval_id"]: db.execute(
            "SELECT * FROM history WHERE evaluation_id = ? ORDER BY changed_at", (a["eval_id"],)).fetchall()
        for a in anomalies
    }


def violated_criteria(db, project_id):
    """Criteri WCAG violati, distinti, per il progetto."""
    return db.execute("""
        SELECT DISTINCT cr.code, cr.title, cr.level, cr.sort_order
        FROM evaluations e
        JOIN components c ON c.id = e.component_id
        JOIN pages pg ON pg.id = c.page_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE pg.project_id = ? AND e.status = 'non_conforme' ORDER BY cr.sort_order""", (project_id,)).fetchall()


def all_evaluations(db, project_id):
    """Tutte le valutazioni del progetto (non solo le anomalie), per l'export XLSX."""
    return db.execute("""
        SELECT pg.name AS page_name, c.name AS component_name, t.name AS type_name,
               cr.code, cr.title, cr.level, e.status, e.work_status, e.note, e.updated_at
        FROM evaluations e
        JOIN components c ON c.id = e.component_id
        JOIN component_types t ON t.id = c.component_type_id
        JOIN pages pg ON pg.id = c.page_id
        JOIN criteria cr ON cr.id = e.criterion_id
        WHERE pg.project_id = ?
        ORDER BY pg.created_at, c.created_at, cr.sort_order""", (project_id,)).fetchall()


def build_csv(project, anomalies):
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Pagina", "URL pagina", "Componente", "Tipologia", "Posizione",
                "Criterio WCAG", "Titolo criterio", "Livello", "Stato lavorazione",
                "Motivazione / Note", "Ultimo aggiornamento"])
    for a in anomalies:
        w.writerow([a["page_name"], a["page_url"], a["component_name"], a["type_name"],
                    a["location"], a["code"], a["title"], a["level"],
                    WORK_LABELS.get(a["work_status"], a["work_status"]),
                    a["note"], a["updated_at"]])
    data = io.BytesIO(("﻿" + out.getvalue()).encode("utf-8"))
    fname = f"anomalie_{project['name'].replace(' ', '_')}.csv" if project else "anomalie.csv"
    return data, fname


def build_xlsx(project, anomalies, all_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")

    def sheet(ws, headers, rows):
        ws.append(headers)
        for i in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=i)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        widths = [len(h) for h in headers]
        for r in rows:
            for i, v in enumerate(r):
                widths[i] = min(60, max(widths[i], len(str(v or ""))))
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt + 2
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws1 = wb.active
    ws1.title = "Anomalie"
    sheet(ws1,
          ["Pagina", "Componente", "Tipologia", "Posizione", "Criterio", "Titolo",
           "Livello", "Stato lavorazione", "Motivazione / Note", "Aggiornato il"],
          [[a["page_name"], a["component_name"], a["type_name"], a["location"],
            a["code"], a["title"], a["level"],
            WORK_LABELS.get(a["work_status"], a["work_status"]), a["note"], a["updated_at"]]
           for a in anomalies])

    ws2 = wb.create_sheet("Tutte le valutazioni")
    sheet(ws2,
          ["Pagina", "Componente", "Tipologia", "Criterio", "Titolo", "Livello",
           "Esito", "Stato lavorazione", "Note", "Aggiornato il"],
          [[r["page_name"], r["component_name"], r["type_name"], r["code"], r["title"],
            r["level"], STATUS_LABELS.get(r["status"], r["status"]),
            WORK_LABELS.get(r["work_status"], r["work_status"]), r["note"], r["updated_at"]]
           for r in all_rows])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"report_{project['name'].replace(' ', '_')}.xlsx" if project else "report.xlsx"
    return buf, fname
